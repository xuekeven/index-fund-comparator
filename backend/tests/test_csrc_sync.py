from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import json
from urllib.parse import unquote

from openpyxl import Workbook

from app.sync.csrc_funds import (
    ScaleRecord,
    classify_product,
    discover_subscription_shares,
    DisclosureDocument,
    disclosure_documents,
    discovery_name,
    extract_index_download_url,
    fetch_nav_records,
    fund_detail_url,
    is_exchange_traded_summary,
    merge_summary_share,
    nav_query_data,
    parse_fund_detail,
    parse_nav_rows,
    parse_product_index,
    parse_product_summary,
    parse_quarterly_scales,
    parse_subscription_announcement,
    resolve_subscription_states,
    share_identity,
    should_retire_stale_catalog_products,
    SummaryShare,
    snapshot_date_from_url,
    subscription_documents,
)


def test_merge_summary_share_prefers_the_more_specific_official_name() -> None:
    shares: dict[str, tuple[SummaryShare, str]] = {}
    merge_summary_share(
        shares,
        SummaryShare("017641", "摩根标普500指数发起式(QDII)", None, "人民币", None),
        "https://example.test/summary.pdf",
    )
    merge_summary_share(
        shares,
        SummaryShare("017641", "摩根标普500指数(QDII)人民币A", "A", "人民币", None),
        "https://example.test/announcement.pdf",
    )

    assert shares["017641"] == (
        SummaryShare("017641", "摩根标普500指数(QDII)人民币A", "A", "人民币", None),
        "https://example.test/announcement.pdf",
    )


def test_stale_catalog_retirement_runs_only_for_complete_full_sync() -> None:
    assert should_retire_stale_catalog_products(failures=[], codes=(), limit=None)
    assert not should_retire_stale_catalog_products(
        failures=[], codes=("050025",), limit=None
    )
    assert not should_retire_stale_catalog_products(failures=[], codes=(), limit=1)
    assert not should_retire_stale_catalog_products(
        failures=["download failed"], codes=(), limit=None
    )


def test_nav_query_uses_the_endpoint_page_limit() -> None:
    query = {
        item["name"]: item["value"]
        for item in json.loads(
            nav_query_data(
                "博时标普500ETF联接A",
                date(2025, 7, 1),
                date(2026, 8, 1),
            )
        )
    }

    assert query["iDisplayStart"] == 0
    assert query["iDisplayLength"] == 500
    assert query["fundCode"] == ""
    assert unquote(query["fundName"]) == "博时标普500ETF联接"
    assert query["startDate"] == "2025-07-01"
    assert query["endDate"] == "2026-08-01"

    by_code = {
        item["name"]: item["value"]
        for item in json.loads(
            nav_query_data(
                "",
                date(2026, 7, 1),
                date(2026, 8, 1),
                fund_code="050025",
            )
        )
    }
    assert by_code["fundName"] == ""
    assert by_code["fundCode"] == "050025"


def test_nav_fetch_splits_a_year_into_supported_date_windows() -> None:
    requests: list[dict[str, str]] = []

    class Response:
        def raise_for_status(self) -> None:
            pass

        def json(self) -> dict[str, object]:
            return {"aaData": []}

    class Client:
        def get(self, _url: str, *, params: dict[str, str]) -> Response:
            query = {
                item["name"]: item["value"]
                for item in json.loads(params["aoData"])
            }
            requests.append(query)
            return Response()

    records = fetch_nav_records(  # type: ignore[arg-type]
        Client(),
        3201,
        "博时标普500ETF联接",
        date(2026, 8, 1),
        lookback_days=400,
    )

    assert records == []
    assert len(requests) == 9
    assert requests[0]["startDate"] == "2025-06-27"
    assert requests[-1]["endDate"] == "2026-08-01"
    for previous, current in zip(requests, requests[1:], strict=False):
        previous_end = date.fromisoformat(previous["endDate"])
        current_start = date.fromisoformat(current["startDate"])
        assert current_start == previous_end + timedelta(days=1)


def test_classifies_target_off_exchange_funds_and_excludes_variants() -> None:
    assert classify_product("嘉实中证500ETF联接")[0].family_id == "csi-500"
    assert classify_product("南方纳斯达克100指数发起（QDII）")[0].family_id == "nasdaq-100"
    assert classify_product("天弘标普500发起（QDII-FOF）")[0].family_id == "sp-500"
    assert classify_product("大成标普500等权重指数(QDII)") is None
    assert classify_product("华夏中证500指数增强") is None
    assert classify_product("博时纳斯达克100ETF") is None
    assert classify_product("鹏华中证500指数（LOF）") is None
    assert is_exchange_traded_summary(
        "上市交易所及上市日期上海证券交易所2023年03月20日"
        "基金类型股票型运作方式交易型开放式开放频率每个开放日"
    )
    assert is_exchange_traded_summary(
        "基金类型股票型运作方式上市开放式开放频率每个开放日"
    )
    assert not is_exchange_traded_summary(
        "上市交易所及上市日期-基金类型基金中基金"
        "运作方式普通开放式开放频率每个开放日"
    )
    assert not is_exchange_traded_summary(
        "基金类型股票型运作方式普通开放式开放频率每个开放日"
    )
    assert not is_exchange_traded_summary(
        "上市交易所及上市日期暂未上市基金类型基金中基金"
        "运作方式开放式(普通开放式）开放频率每个开放日"
    )
    assert not is_exchange_traded_summary(
        "上市交易所及上市日期（若有）-基金类型股票型"
        "运作方式普通开放式开放频率每个开放日"
    )


def test_parses_csrc_product_index() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["公募基金产品索引（截至20260630）"])
    sheet.append(["序号", "基金代码", "基金简称", "设立日期"])
    sheet.append([1, 8, "嘉实中证500ETF联接", "2013-03-22"])
    sheet.append([2, "510500", "中证500ETF", "2013-02-06"])
    sheet.append([3, "513110", "纳指100", "2022-04-13"])
    sheet.append([4, "160213", "国泰纳斯达克100指数（QDII）", "2010-04-29"])
    sheet.append([5, "164809", "工银中证500ETF联接", "2020-02-18"])
    output = BytesIO()
    workbook.save(output)

    candidates, snapshot_date = parse_product_index(output.getvalue())

    assert snapshot_date == date(2026, 6, 30)
    assert [(item.code, item.product_structure) for item in candidates] == [
        ("000008", "ETF联接基金"),
        ("513110", "普通开放式指数基金"),
        ("160213", "普通开放式指数基金"),
        ("164809", "ETF联接基金"),
    ]


def test_extracts_relative_index_download_url() -> None:
    page = '<a href="1029655/files/public-funds.xlsx">download</a>'
    assert extract_index_download_url(page) == (
        "https://www.csrc.gov.cn/csrc/c101900/c1029655/"
        "1029655/files/public-funds.xlsx"
    )
    assert snapshot_date_from_url(
        "https://example.test/公募基金产品索引（截至20260630）.xlsx"
    ) == date(2026, 6, 30)


def test_parses_eid_detail_page() -> None:
    page = """
    <td class="title_tu">博时标普500ETF联接(050025)</td>
    <td>基金名称</td><td>博时标普500交易型开放式指数证券投资基金联接基金</td>
    <td>基金代码</td><td>050025</td>
    <td>基金管理人</td><td>博时基金管理有限公司</td>
    <td>基金合同生效日期</td><td>2012-06-14</td>
    """
    detail = parse_fund_detail(3201, page)
    assert detail.short_name == "博时标普500ETF联接"
    assert detail.manager == "博时基金管理有限公司"
    assert detail.inception_date == date(2012, 6, 14)
    assert fund_detail_url(detail.fund_id) == (
        "http://eid.csrc.gov.cn/fund/disclose/fund_detail.do?fundId=3201"
    )


def test_discovers_share_identity_and_filters_product_level_nav_rows() -> None:
    rows = [
        {
            "code": "050025",
            "shortName": "博时标普500ETF联接A",
            "valuationDate": "2026-08-12",
            "shareNetValue": "5.6433",
            "totalNetValue": "5.7023",
            "fund": {"idStr": 3201},
        },
        {
            "code": "006075",
            "shortName": "博时标普500ETF联接C",
            "valuationDate": "2026-08-12",
            "shareNetValue": "5.1234",
            "totalNetValue": "5.1234",
            "fund": {"idStr": 3201},
        },
        {
            "code": "050025",
            "shortName": "博时标普500ETF联接",
            "valuationDate": "2026-08-12",
            "shareNetValue": "",
            "fund": {"idStr": 3201},
        },
    ]
    records = parse_nav_rows(rows, 3201)
    assert [(item.code, item.share_class) for item in records] == [
        ("006075", "C"),
        ("050025", "A"),
    ]
    assert discovery_name("建信纳斯达克100指数(QDII)A人民币") == (
        "建信纳斯达克100指数(QDII)"
    )
    assert share_identity("某基金美元现汇A") == ("A", "美元", "现汇")
    assert share_identity("某基金美汇") == (None, "美元", "现汇")
    assert share_identity("某基金美钞") == (None, "美元", "现钞")


def test_extracts_disclosure_documents_by_type_and_deduplicates() -> None:
    page = """
    <a href="../instance_show_pdf_id.do?instanceid=12">基金产品资料概要更新</a>
    <a href="../instance_show_pdf_id.do?instanceid=12">基金产品资料概要更新</a>
    <a href="../instance_show_pdf_id.do?instanceid=34">2026年第2季度报告</a>
    """

    summaries = disclosure_documents(page, "基金产品资料概要")
    reports = disclosure_documents(page, "季度报告")

    assert len(summaries) == 1
    assert summaries[0].url == (
        "http://eid.csrc.gov.cn/fund/instance_show_pdf_id.do?instanceid=12"
    )
    assert len(reports) == 1
    assert reports[0].title == "2026年第2季度报告"


def test_parses_product_summary_fees_codes_and_benchmark() -> None:
    text = """
    基金简称 博时标普500ETF联接 基金代码 050025
    下属基金简称 A C 下属基金代码 050025 013499
    业绩比较基准 标普500指数收益率（经估值汇率调整）×95%+
    银行活期存款利率（税后）×5% 风险收益特征 本基金属于股票型基金
    管理费 固定比例 0.60%
    托管费 固定比例 0.20%
    销售服务费 固定比例 0.35%
    """

    summary = parse_product_summary(text)

    assert summary.share_codes == ("050025", "013499")
    assert summary.rates == {
        "management": Decimal("0.60"),
        "custody": Decimal("0.20"),
        "sales_service": Decimal("0.35"),
    }
    assert summary.benchmark_description == (
        "标普500指数收益率（经估值汇率调整）×95%+"
        "银行活期存款利率（税后）×5%"
    )


def test_product_summary_identifies_050025_as_the_a_share() -> None:
    summary = parse_product_summary(
        "博时标普500ETF联接A基金产品资料概要更新"
        "基金简称博时标普500ETF联接A基金代码050025"
        "管理费0.60%托管费0.20%"
    )

    assert summary.shares == (
        SummaryShare(
            "050025",
            "博时标普500ETF联接A",
            "A",
            "人民币",
            None,
        ),
    )


def test_product_summary_defaults_missing_sales_service_to_zero() -> None:
    summary = parse_product_summary(
        "基金代码000008业绩比较基准中证500指数收益率×95%"
        "+银行活期存款利率×5%风险收益特征较高风险"
        "管理费0.50%托管费0.10%"
    )
    assert summary.rates["sales_service"] == Decimal("0")



def test_product_summary_accepts_subordinate_trading_code() -> None:
    summary = parse_product_summary(
        "基金代码016532下属基金交易代码021838"
        "业绩比较基准纳斯达克指数收益率风险收益特征较高风险"
        "管理费0.80%托管费0.20%销售服务费0.10%"
    )

    assert summary.share_codes == ("021838",)
    assert summary.rates["sales_service"] == Decimal("0.10")


def test_product_summary_discovers_usd_cash_share() -> None:
    summary = parse_product_summary(
        "博时标普500ETF联接A（美元现汇）基金产品资料概要更新"
        "基金简称博时标普500ETF联接基金代码050025"
        "下属基金简称博时标普500ETF联接A（美元现汇）"
        "下属基金交易代码013425"
        "管理费0.60%托管费0.20%"
    )

    assert summary.shares == (
        SummaryShare(
            "013425",
            "博时标普500ETF联接A（美元现汇）",
            "A",
            "美元",
            "现汇",
        ),
    )


def test_parses_subscription_suspension_flags_and_limits() -> None:
    shares = {
        share.code: share
        for share in (
            SummaryShare("013425", "博时标普500ETF联接A（美元现汇）", "A", "美元", "现汇"),
            SummaryShare("050025", "博时标普500ETF联接A", "A", "人民币", None),
            SummaryShare("006075", "博时标普500ETF联接C", "C", "人民币", None),
            SummaryShare("013499", "博时标普500ETF联接C（美元现汇）", "C", "美元", "现汇"),
            SummaryShare("018738", "博时标普500ETF联接E", "E", "人民币", None),
        )
    }
    suspension = parse_subscription_announcement(
        "A、C两类份额暂停申购业务公告",
        "公告送出日期2026年5月8日暂停申购起始日2026年5月11日"
        "下属分级基金的交易代码013425050025006075013499018738"
        "该分级基金是否暂停申购是是是是否2、其他需要提示的事项",
        shares,
        "https://example.test/suspended.pdf",
    )
    assert [(item.code, item.status, item.effective_date) for item in suspension] == [
        ("013425", "suspended", date(2026, 5, 11)),
        ("050025", "suspended", date(2026, 5, 11)),
        ("006075", "suspended", date(2026, 5, 11)),
        ("013499", "suspended", date(2026, 5, 11)),
    ]

    limited = parse_subscription_announcement(
        "调整大额申购业务公告",
        "公告送出日期2026年4月9日暂停大额申购起始日2026年4月10日"
        "下属分级基金的交易代码013425050025006075013499018738"
        "下属分级基金是否暂停大额申购是是是是是2、其他需要提示的事项"
        "申购本基金A类人民币份额单日累计金额应不超过100元人民币，"
        "申购本基金C类人民币份额单日累计金额应不超过100元人民币，"
        "申购本基金E类人民币份额单日累计金额应不超过2000元人民币。"
        "申购本基金A类美元现汇份额单日累计金额应不超过15美元，"
        "申购本基金C类美元现汇份额单日累计金额应不超过15美元。",
        shares,
        "https://example.test/limited.pdf",
    )
    assert {item.code: item.limit_amount for item in limited} == {
        "013425": Decimal("15"),
        "050025": Decimal("100"),
        "006075": Decimal("100"),
        "013499": Decimal("15"),
        "018738": Decimal("2000"),
    }


def test_parses_platform_suspension_when_fund_name_sits_between_keywords() -> None:
    shares = {
        share.code: share
        for share in (
            SummaryShare("018064", "华夏标普500ETF发起式联接A（人民币）", "A", "人民币", None),
            SummaryShare("018065", "华夏标普500ETF发起式联接C（人民币）", "C", "人民币", None),
            SummaryShare("018066", "华夏标普500ETF发起式联接A（美元现汇）", "A", "美元", "现汇"),
        )
    }

    states = parse_subscription_announcement(
        "关于在基金管理人直销电子交易平台暂停华夏标普500ETF发起式联接基金（QDII）人民币申购业务的公告",
        "公告送出日期2026年8月7日，自2026年8月10日起，"
        "涉及基金份额类别的交易代码018064018065"
        "基金管理人直销电子交易平台暂停人民币申购业务。",
        shares,
        "http://eid.csrc.gov.cn/fund/disclose/instance_show_pdf_id.do?instanceid=1548363",
    )

    assert [(state.code, state.status, state.channel) for state in states] == [
        ("018064", "suspended", "基金管理人直销电子交易平台"),
        ("018065", "suspended", "基金管理人直销电子交易平台"),
    ]


def test_preserves_pdf_table_columns_for_equal_subscription_limits() -> None:
    shares = {
        share.code: share
        for share in (
            SummaryShare("018064", "华夏标普500ETF发起式联接A（人民币）", "A", "人民币", None),
            SummaryShare("018065", "华夏标普500ETF发起式联接C（人民币）", "C", "人民币", None),
        )
    }

    states = parse_subscription_announcement(
        "恢复申购业务并限制申购金额的公告",
        """公告送出日期2026年6月26日，自2026年6月29日起
涉及基金份额类别的交易代
码 018064 018065
该基金份额类别在基金管理
人直销电子交易平台限制申
购金额（单位：人民币元）
100 100
注：其他事项
""",
        shares,
        "http://eid.csrc.gov.cn/fund/disclose/instance_show_pdf_id.do?instanceid=1513627",
    )

    assert [(state.code, state.limit_amount) for state in states] == [
        ("018064", Decimal("100")),
        ("018065", Decimal("100")),
    ]


def test_resolves_subscription_state_by_effective_date_not_document_order() -> None:
    shares = {
        "018064": SummaryShare(
            "018064", "华夏标普500ETF发起式联接A（人民币）", "A", "人民币", None
        )
    }
    announcements = [
        (
            DisclosureDocument("恢复申购并限制金额公告", "https://example.test/old.pdf"),
            "自2026年6月29日起限制申购金额（单位：人民币元）100元",
        ),
        (
            DisclosureDocument(
                "暂停华夏标普500ETF发起式联接基金人民币申购业务的公告",
                "https://example.test/new.pdf",
            ),
            "自2026年8月10日起暂停人民币申购业务",
        ),
    ]

    states, warnings = resolve_subscription_states(
        announcements, shares, date(2026, 8, 30), "https://example.test/detail"
    )

    assert warnings == []
    assert states["018064"].status == "suspended"
    assert states["018064"].effective_date == date(2026, 8, 10)


def test_parses_combined_share_limit_in_yuan_and_ignores_dash_flag() -> None:
    shares = {
        share.code: share
        for share in (
            SummaryShare("270042", "某基金A", "A", "人民币", None),
            SummaryShare("006479", "某基金C", "C", "人民币", None),
            SummaryShare("021778", "某基金F", "F", "人民币", None),
            SummaryShare("000055", "某基金A美元", "A", "美元", None),
            SummaryShare("006480", "某基金C美元", "C", "美元", None),
        )
    }

    states = parse_subscription_announcement(
        "暂停大额申购业务公告",
        "暂停大额申购起始日2026年6月1日"
        "下属基金份额的交易代码270042006479021778000055006480"
        "下属分级基金是否暂停大额申购是是-是是2、其他需要提示的事项"
        "申购本基金A类及C类人民币份额业务限额为5.00元，"
        "申购本基金A类及C类美元份额业务限额为20.00美元。",
        shares,
        "https://example.test/limited.pdf",
    )

    assert [(item.code, item.limit_amount) for item in states] == [
        ("270042", Decimal("5.00")),
        ("006479", Decimal("5.00")),
        ("000055", Decimal("20.00")),
        ("006480", Decimal("20.00")),
    ]


def test_parses_amount_limit_announcements_by_share_class() -> None:
    shares = {
        share.code: share
        for share in (
            SummaryShare("016452", "南方纳指100A", "A", "人民币", None),
            SummaryShare("016453", "南方纳指100C", "C", "人民币", None),
            SummaryShare("021000", "南方纳指100I", "I", "人民币", None),
        )
    }
    announcements = [
        (
            DisclosureDocument(
                "关于调整南方纳斯达克100指数发起式证券投资基金（QDII）I类基金份额申购、定投及转换转入业务金额限制的公告",
                "https://example.test/i.pdf",
            ),
            "公告送出日期：2026年7月20日本公司自2026年7月21日起调整"
            "下属基金份额的交易代码016452016453021000"
            "该基金份额的限制金额200元本次调整不涉及本基金A类、C类基金份额。",
        ),
        (
            DisclosureDocument(
                "关于调整南方纳斯达克100指数发起式证券投资基金（QDII）申购、定投及转换转入业务金额限制的公告",
                "https://example.test/ac.pdf",
            ),
            "公告送出日期：2026年7月8日本公司自2026年7月9日起调整"
            "下属基金份额的交易代码016452016453021000"
            "调整后本基金A类、C类基金份额限额10元，I类基金份额仍保持1000元限额不变。",
        ),
    ]

    states, warnings = resolve_subscription_states(
        announcements,
        shares,
        date(2026, 8, 28),
        "https://example.test/detail",
    )

    assert warnings == []
    assert {
        code: (state.status, state.limit_amount, state.effective_date)
        for code, state in states.items()
    } == {
        "016452": ("limited", Decimal("10"), date(2026, 7, 9)),
        "016453": ("limited", Decimal("10"), date(2026, 7, 9)),
        "021000": ("limited", Decimal("200"), date(2026, 7, 21)),
    }


def test_restore_subscription_with_large_limit_stays_limited() -> None:
    shares = {
        "160213": SummaryShare(
            "160213",
            "国泰纳斯达克100指数（QDII）",
            None,
            "人民币",
            None,
        )
    }

    states = parse_subscription_announcement(
        "关于国泰纳斯达克100指数证券投资基金恢复申购及定期定额投资业务（限制大额）的公告",
        "公告送出日期2026年8月18日本基金管理人决定自2026年8月19日起"
        "恢复本基金的申购及定期定额投资业务（限制大额）。"
        "单个基金账户对本基金单日累计金额50.00元以下（含50.00元）的"
        "申购及定期定额投资业务申请，如累计金额超过50.00元，"
        "本基金管理人有权部分或全部确认失败。",
        shares,
        "https://example.test/160213.pdf",
    )

    assert [
        (state.code, state.status, state.limit_amount, state.effective_date)
        for state in states
    ] == [("160213", "limited", Decimal("50.00"), date(2026, 8, 19))]


def test_subscription_documents_include_amount_limit_titles() -> None:
    html = """
    <a href="instance_show_pdf_id.do?instanceid=1">申购、定投业务金额限制公告</a>
    <a href="instance_show_pdf_id.do?instanceid=2">恢复大额申购公告</a>
    <a href="instance_show_pdf_id.do?instanceid=3">境外节假日申购安排公告</a>
    <a href="instance_show_pdf_id.do?instanceid=4">基金产品资料概要</a>
    """

    assert [document.title for document in subscription_documents(html)] == [
        "申购、定投业务金额限制公告",
        "恢复大额申购公告",
    ]


def test_unknown_subscription_state_is_not_assumed_open() -> None:
    shares = {
        "016452": SummaryShare("016452", "南方纳指100A", "A", "人民币", None)
    }

    states, warnings = resolve_subscription_states(
        [],
        shares,
        date(2026, 8, 28),
        "https://example.test/detail",
    )

    assert states == {}
    assert warnings == []


def test_discovers_usd_cash_share_from_subscription_announcement() -> None:
    document = DisclosureDocument("暂停大额申购公告", "https://example.test/a.pdf")
    shares = discover_subscription_shares(
        [
            (
                document,
                "下属分级基金的交易代码040046040047014978040048"
                "分别为A类人民币基金份额（基金代码：040046），"
                "A类美元现钞基金份额（基金代码：040047），"
                "C类人民币基金份额（基金代码：014978），"
                "A类美元现汇基金份额（基金代码：040048）。",
            )
        ],
        "华安纳斯达克100ETF联接（QDII）",
    )

    assert shares["040048"][0] == SummaryShare(
        "040048",
        "华安纳斯达克100ETF联接（QDII）A（美元现汇）",
        "A",
        "美元",
        "现汇",
    )


def test_discovers_exact_share_names_from_subscription_announcement() -> None:
    document = DisclosureDocument("暂停申购公告", "https://example.test/a.pdf")
    shares = discover_subscription_shares(
        [
            (
                document,
                "下属分级基金的基金简称"
                "博时标普500ETF联接A（人民币）"
                "博时标普500ETF联接A（美元现汇）"
                "下属分级基金的交易代码050025013425"
                "金额单位人民币元美元下属基金份额的限制申购金额",
            )
        ],
        "博时标普500ETF联接",
    )

    assert shares["050025"][0].display_name == "博时标普500ETF联接A（人民币）"
    assert shares["013425"][0].display_name == "博时标普500ETF联接A（美元现汇）"


def test_discovers_richer_names_after_a_less_specific_announcement() -> None:
    generic = DisclosureDocument("暂停申购公告", "https://example.test/old.pdf")
    exact = DisclosureDocument("调整大额申购公告", "https://example.test/new.pdf")
    shares = discover_subscription_shares(
        [
            (
                generic,
                "下属分级基金的简称摩根标普500指数(QDII)"
                "下属分级基金的交易代码017641019305",
            ),
            (
                exact,
                "下属分级基金的基金简称"
                "摩根标普500指数(QDII)人民币A"
                "摩根标普500指数(QDII)人民币C"
                "下属分级基金的交易代码017641019305",
            ),
        ],
        "摩根标普500指数发起式(QDII)",
    )

    assert shares["017641"][0] == SummaryShare(
        "017641", "摩根标普500指数(QDII)人民币A", "A", "人民币", None
    )
    assert shares["019305"][0] == SummaryShare(
        "019305", "摩根标普500指数(QDII)人民币C", "C", "人民币", None
    )


def test_parses_quarterly_share_scales() -> None:
    text = """
    博时标普500ETF联接2026年第2季度报告
    下属分级基金的交易代码 050025 006075 018738
    4.期末基金资产净值
    6,466,845,879.52 1,285,578,443.52 2,304,202,372.31
    5.期末基金份额净值 5.1 4.9 1.2
    """

    scales = parse_quarterly_scales(text)

    assert [item.share_code for item in scales] == [
        "050025",
        "006075",
        "018738",
    ]
    assert {item.report_date for item in scales} == {date(2026, 6, 30)}
    assert scales[0].amount_cny == Decimal("6466845879.52")


def test_parses_fees_after_long_etf_deduction_description() -> None:
    summary = parse_product_summary(
        "基金代码001052业绩比较基准中证500指数收益率风险收益特征较高风险"
        "管理费按前一日基金资产净值扣除基金资产中目标ETF份额所对应资产"
        "净值后剩余部分的年费率计提，并按基金合同约定每日计算和支付。0.15%"
        "托管费按前一日基金资产净值扣除目标ETF对应资产后计提。0.05%"
    )

    assert summary.rates["management"] == Decimal("0.15")
    assert summary.rates["custody"] == Decimal("0.05")


def test_parses_single_share_report_with_main_code_and_chinese_year() -> None:
    scales = parse_quarterly_scales(
        "某基金二0二六年第2季度报告基金主代码001241"
        "4.期末基金资产净值133,322,757.45"
        "5.期末基金份额净值0.9895"
    )

    assert scales == [
        ScaleRecord("001241", date(2026, 6, 30), Decimal("133322757.45"))
    ]


def test_parses_scale_when_next_heading_is_split_by_page_text() -> None:
    scales = parse_quarterly_scales(
        "某基金2026年第2季度报告"
        "下属分级基金的交易代码016532016533021838"
        "4.期末基金资产净值"
        "2,076,998,722.981,640,123,827.7824,300,350.52"
        "5.期末基金份2.23332.20902.2139额净值"
    )

    assert [item.amount_cny for item in scales] == [
        Decimal("2076998722.98"),
        Decimal("1640123827.78"),
        Decimal("24300350.52"),
    ]
